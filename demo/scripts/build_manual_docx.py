#!/usr/bin/env python
# -*- coding: UTF-8 -*-
'''
将答辩镜像操作手册 Markdown 转成可打印的 Word 文档。
@Project : adaptive-edu
@File : build_manual_docx.py
@Author : Qintsg
@Date : 2026-09-27
'''

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor
from docx.styles.style import ParagraphStyle
from docx.table import Table, _Cell as Cell
from docx.text.paragraph import Paragraph
from docx.text.run import Run


FONT_NAME = "Microsoft YaHei"
CODE_FONT_NAME = "Consolas"
INK = RGBColor(20, 34, 50)
SUBTLE = RGBColor(65, 77, 91)
HEADER_FILL = "18324F"
BODY_FILL = "F3F7FB"
BORDER_COLOR = "D9D9D9"


def set_font(run: Run, name: str, size_pt: float, color: RGBColor,
             bold: bool = False) -> None:
    """设置中英文字体，避免 Word 与渲染器使用不同的回退字体。

    :param run: 要格式化的文字片段。
    :param name: 字体名称。
    :param size_pt: 字号。
    :param color: 字体颜色。
    :param bold: 是否加粗。
    :returns: 无。
    """
    run.font.name = name
    run.font.size = Pt(size_pt)
    run.font.color.rgb = color
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attr}"), name)


def set_style_font(style: ParagraphStyle, name: str, size_pt: float,
                   color: RGBColor, bold: bool = False) -> None:
    """为段落样式设置字体及东亚字体。

    :param style: 段落样式。
    :param name: 字体名称。
    :param size_pt: 字号。
    :param color: 字体颜色。
    :param bold: 是否加粗。
    :returns: 无。
    """
    style.font.name = name
    style.font.size = Pt(size_pt)
    style.font.color.rgb = color
    style.font.bold = bold
    rpr = style._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attr}"), name)


def configure_document(document: DocumentType) -> None:
    """设置 Letter 纵向页面和便于现场查阅的文字层级。

    :param document: 目标 Word 文档。
    :returns: 无。
    """
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Mm(21)
    section.right_margin = Mm(21)
    section.top_margin = Mm(17)
    section.bottom_margin = Mm(16)

    normal = document.styles["Normal"]
    set_style_font(normal, FONT_NAME, 10.5, INK)
    normal.paragraph_format.line_spacing = 1.24
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.keep_together = True

    title = document.styles["Title"]
    set_style_font(title, FONT_NAME, 20, RGBColor(0, 0, 0), True)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(11)
    title.paragraph_format.keep_with_next = True
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    heading = document.styles["Heading 1"]
    set_style_font(heading, FONT_NAME, 13.5, RGBColor(0, 0, 0), True)
    heading.paragraph_format.space_before = Pt(13)
    heading.paragraph_format.space_after = Pt(6)
    heading.paragraph_format.keep_with_next = True

    for list_name in ("List Bullet", "List Number"):
        list_style = document.styles[list_name]
        set_style_font(list_style, FONT_NAME, 10.5, INK)
        list_style.paragraph_format.left_indent = Mm(7)
        list_style.paragraph_format.first_line_indent = Mm(-3)
        list_style.paragraph_format.space_after = Pt(4)
        list_style.paragraph_format.line_spacing = 1.18

    code_style = document.styles.add_style("Manual Code", 1)
    set_style_font(code_style, CODE_FONT_NAME, 8.7, SUBTLE)
    code_style.paragraph_format.left_indent = Mm(4)
    code_style.paragraph_format.space_after = Pt(2)
    code_style.paragraph_format.line_spacing = 1.12

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_before = Pt(0)
    label = footer.add_run("第 ")
    set_font(label, FONT_NAME, 8.5, SUBTLE)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    suffix = footer.add_run(" 页")
    set_font(suffix, FONT_NAME, 8.5, SUBTLE)


def add_inline(paragraph: Paragraph, text: str) -> None:
    """添加普通文本、反引号等宽文本和双星号加粗文本。

    :param paragraph: 目标段落。
    :param text: Markdown 行内容。
    :returns: 无。
    """
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        code = part.startswith("`") and part.endswith("`")
        bold = part.startswith("**") and part.endswith("**")
        value = part[1:-1] if code else part[2:-2] if bold else part
        run = paragraph.add_run(value)
        if code:
            set_font(run, CODE_FONT_NAME, 9, INK)
        elif bold:
            run.bold = True


def cell_margin(cell: Cell, top: int = 95, start: int = 105,
                bottom: int = 95, end: int = 105) -> None:
    """设置表格单元格留白，单位为 twip。

    :param cell: 目标单元格。
    :param top: 上边距。
    :param start: 左边距。
    :param bottom: 下边距。
    :param end: 右边距。
    :returns: 无。
    """
    tcpr = cell._tc.get_or_add_tcPr()
    margins = tcpr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tcpr.append(margins)
    for side, size in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(size))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell: Cell, fill: str) -> None:
    """设置单元格底色。

    :param cell: 目标单元格。
    :param fill: RGB 十六进制颜色。
    :returns: 无。
    """
    tcpr = cell._tc.get_or_add_tcPr()
    shading = tcpr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tcpr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_table_borders(table: Table) -> None:
    """给表格添加浅灰色外框和内部边框。

    :param table: 目标表格。
    :returns: 无。
    """
    tblpr = table._tbl.tblPr
    borders = tblpr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblpr.append(borders)
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = borders.find(qn(f"w:{side}"))
        if border is None:
            border = OxmlElement(f"w:{side}")
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "5")
        border.set(qn("w:color"), BORDER_COLOR)


def set_row_rules(row: object, repeat_header: bool = False) -> None:
    """让表头跨页重复，并避免单行被拆到两页。

    :param row: python-docx 表格行。
    :param repeat_header: 是否为表头。
    :returns: 无。
    """
    trpr = row._tr.get_or_add_trPr()
    no_split = OxmlElement("w:cantSplit")
    trpr.append(no_split)
    if repeat_header:
        header = OxmlElement("w:tblHeader")
        header.set(qn("w:val"), "true")
        trpr.append(header)


def parse_table_line(line: str) -> list[str]:
    """拆分 Markdown 表格行。

    :param line: 原始表格行。
    :returns: 清洗后的单元格内容。
    """
    return [value.strip() for value in line.strip().strip("|").split("|")]


def add_table(document: DocumentType, header: list[str], rows: list[list[str]]) -> None:
    """用固定列宽写入表格，并维持跨页可读性。

    :param document: 目标文档。
    :param header: 表头内容。
    :param rows: 表体内容。
    :returns: 无。
    """
    widths = [Mm(48), Mm(125)] if len(header) == 2 else [Mm(28), Mm(26), Mm(28), Mm(91)]
    table = document.add_table(rows=1, cols=len(header))
    table.autofit = False
    set_table_borders(table)
    for index, value in enumerate(header):
        cell = table.rows[0].cells[index]
        cell.width = widths[index]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell_margin(cell)
        shade_cell(cell, HEADER_FILL)
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.13
        run = paragraph.add_run(value)
        set_font(run, FONT_NAME, 9.2, RGBColor(255, 255, 255), True)
    set_row_rules(table.rows[0], True)

    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        set_row_rules(table.rows[-1])
        for column_index, value in enumerate(values):
            cell = cells[column_index]
            cell.width = widths[column_index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_margin(cell)
            if row_index % 2:
                shade_cell(cell, BODY_FILL)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.16
            add_inline(paragraph, value)
            for run in paragraph.runs:
                if run.font.name != CODE_FONT_NAME:
                    set_font(run, FONT_NAME, 9.2, INK)
            if column_index in (0, 1, 2) and len(header) == 4:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if len(header) == 4:
        for row in table.rows[:-1]:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.keep_with_next = True
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def render_markdown(document: DocumentType, source: str) -> None:
    """把手册中的标题、段落、表格、列表和命令写入 Word。

    :param document: 目标文档。
    :param source: Markdown 原文。
    :returns: 无。
    """
    lines = source.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line.startswith("```"):
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                paragraph = document.add_paragraph(style="Manual Code")
                paragraph.paragraph_format.keep_together = True
                paragraph.add_run(lines[index])
                index += 1
            index += 1
            continue
        if line.startswith("|") and index + 1 < len(lines) and re.match(r"^\|[\s|:-]+\|$", lines[index + 1].strip()):
            header = parse_table_line(line)
            rows: list[list[str]] = []
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(parse_table_line(lines[index]))
                index += 1
            add_table(document, header, rows)
            continue
        if line.startswith("# "):
            paragraph = document.add_paragraph(style="Title")
            add_inline(paragraph, line[2:])
        elif line.startswith("## "):
            paragraph = document.add_paragraph(style="Heading 1")
            add_inline(paragraph, line[3:])
        elif re.match(r"^\d+\.\s+", line):
            paragraph = document.add_paragraph(style="List Number")
            add_inline(paragraph, re.sub(r"^\d+\.\s+", "", line))
        elif line.startswith("- "):
            paragraph = document.add_paragraph(style="Normal")
            paragraph.paragraph_format.left_indent = Mm(7)
            paragraph.paragraph_format.first_line_indent = Mm(-4)
            bullet = paragraph.add_run("•  ")
            set_font(bullet, FONT_NAME, 10.5, INK)
            add_inline(paragraph, line[2:])
        else:
            paragraph = document.add_paragraph(style="Normal")
            add_inline(paragraph, line)
        index += 1


def main() -> None:
    """读取 Markdown 手册并生成 DOCX。

    :returns: 无。
    """
    parser = argparse.ArgumentParser(description="从 Markdown 更新答辩镜像 Word 操作手册")
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    document = Document()
    configure_document(document)
    render_markdown(document, args.source.read_text(encoding="utf-8"))
    args.output.parent.mkdir(parents=True, exist_ok=True)
    document.save(args.output)
    print(args.output)


if __name__ == "__main__":
    main()
